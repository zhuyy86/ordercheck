"""Batch I/O and auditable exports for the Dataset Cleaner OrderCheck extension."""
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
import csv
import hashlib
import json
import zipfile

import pandas as pd
from openpyxl import load_workbook

from data_cleaning import DataCleaner
from order_rules import FIELDS, PROVENANCE, validate_config
from utils import get_dataset_info

MAX_BYTES = 5 * 1024 * 1024
MAX_EXPANDED = 25 * 1024 * 1024
MAX_ROWS = 20000
MAX_FILES = 20


@dataclass
class InputTable:
    filename: str
    sheet: str
    headers: list
    rows: list
    sha256: str

    @property
    def key(self):
        return f"{self.filename} / {self.sheet}"


def _text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _table(filename, sheet, rows, digest):
    rows = list(rows)
    if not rows:
        raise ValueError("空表：没有表头")
    headers = [_text(x).strip() for x in rows[0]]
    if (not headers or len(headers) > 100 or any(not h for h in headers)
            or len({h.casefold() for h in headers}) != len(headers)):
        raise ValueError("表头为空、重复或超过 100 列")
    data = []
    for number, row in enumerate(rows[1:], start=2):
        if len(row) != len(headers):
            raise ValueError(f"记录 {number} 的列数与表头不一致")
        values = [_text(v) for v in row]
        if any(any(ord(c) < 32 and c not in "\t\n\r" for c in v) for v in values):
            raise ValueError(f"记录 {number} 含不支持的控制字符")
        if len(json.dumps(dict(zip(headers, values)), ensure_ascii=False)) > 30000:
            raise ValueError(f"记录 {number} 过长，无法完整放入审计 Excel")
        data.append((number, values))
        if len(data) > MAX_ROWS:
            raise ValueError(f"单批最多 {MAX_ROWS} 条记录")
    if not data:
        raise ValueError("只有表头，没有数据")
    return InputTable(filename, sheet, headers, data, digest)


def read_tables(filename, content):
    """Read strictly; never let malformed records disappear via skip-bad-lines."""
    if not content or len(content) > MAX_BYTES:
        raise ValueError(f"{filename}: 文件为空或超过 5 MB")
    filename = Path(filename).name
    suffix = Path(filename).suffix.lower()
    digest = hashlib.sha256(content).hexdigest()
    try:
        if suffix == ".csv":
            decoded = content.decode("utf-8-sig")
            sample = decoded[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                reader = csv.reader(StringIO(decoded, newline=""), dialect, strict=True)
            except csv.Error:
                reader = csv.reader(StringIO(decoded, newline=""), delimiter=",", strict=True)
            # Bound rows before materializing an untrusted input.
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError(f"单批最多 {MAX_ROWS} 条记录")
            return [_table(filename, "CSV", rows, digest)]
        if suffix != ".xlsx":
            raise ValueError("仅支持 UTF-8 CSV 和 .xlsx（不支持旧 .xls）")
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED:
                raise ValueError("Excel 解压后超过 25 MB")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            if len(book.worksheets) > 20:
                raise ValueError("Excel 最多 20 个工作表")
            tables = []
            for sheet in book.worksheets:
                if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > 100:
                    raise ValueError(f"{sheet.title}: 工作表行列过多")
                rows = []
                for row in sheet.iter_rows():
                    if any(c.data_type == "f" for c in row):
                        raise ValueError(f"{sheet.title}: 包含公式，请先另存为值")
                    rows.append([c.value for c in row])
                # Empty sheets explicitly rejected rather than silently skipped.
                tables.append(_table(filename, sheet.title, rows, digest))
            return tables
        finally:
            book.close()
    except (ValueError, UnicodeError, csv.Error, zipfile.BadZipFile, OSError, KeyError) as exc:
        raise ValueError(f"{filename}: {exc}") from exc
    except Exception as exc:
        raise ValueError(f"{filename}: 不能读取此文件（{type(exc).__name__}）") from exc


def suggested_mapping(table, config):
    mapping = {}
    for field in FIELDS:
        aliases = {x.strip().casefold() for x in [field, *config["aliases"][field]]}
        matches = [h for h in table.headers if h.casefold() in aliases]
        mapping[field] = matches[0] if len(matches) == 1 else ""
    return mapping


def mapped_frame(table, config, mapping):
    if not isinstance(mapping, dict) or set(mapping) != set(FIELDS):
        raise ValueError(f"{table.key}: 映射必须包含四个标准字段")
    picked = []
    for field, header in mapping.items():
        if not isinstance(header, str):
            raise ValueError("列名映射必须是字符串")
        if header and header not in table.headers:
            raise ValueError(f"{table.key}: 找不到列 {header}")
        if not header and field in config["required"]:
            raise ValueError(f"{table.key}: 请为 {field} 选择来源列（缺失或歧义）")
        if header:
            picked.append(header)
    if len(picked) != len(set(picked)):
        raise ValueError(f"{table.key}: 同一来源列不能映射到多个字段")
    records = []
    for row_number, values in table.rows:
        raw = dict(zip(table.headers, values))
        record = {field: raw.get(mapping[field], "") for field in FIELDS}
        record.update(source_file=table.filename, source_sheet=table.sheet,
                      source_row=str(row_number),
                      raw_json=json.dumps(raw, ensure_ascii=False))
        records.append(record)
    return pd.DataFrame(records, columns=[*FIELDS, *PROVENANCE])


def run_batch(tables, config, mappings=None):
    validate_config(config)
    if not tables:
        raise ValueError("请先提供文件")
    if len({t.filename for t in tables}) > MAX_FILES or len(tables) > 50:
        raise ValueError("单批最多 20 个文件 / 50 个工作表")
    if len({t.key for t in tables}) != len(tables):
        raise ValueError("文件名/工作表重复，请先重命名文件")
    if sum(len(t.rows) for t in tables) > MAX_ROWS:
        raise ValueError(f"单批最多 {MAX_ROWS} 条记录")
    overrides = mappings or {}
    if not isinstance(overrides, dict) or set(overrides) - {t.key for t in tables}:
        raise ValueError("映射配置包含未知文件或工作表")
    resolved, frames = {}, []
    for table in tables:
        mapping = overrides.get(table.key, suggested_mapping(table, config))
        frames.append(mapped_frame(table, config, mapping))
        resolved[table.key] = mapping
    combined = pd.concat(frames, ignore_index=True)
    # Use the original project's DataCleaner, extended with a strict order mode.
    cleaner = DataCleaner(combined)
    result = cleaner.validate_order_batch(config)
    result.summary["files"] = len({t.filename for t in tables})
    result.summary["tables"] = len(tables)
    result.summary["upstream_rows_after"] = cleaner.get_transformations_summary()["rows_after"]
    result.summary["profile_columns"] = get_dataset_info(cleaner.get_cleaned_data())["columns"]
    manifest = {
        "inputs": [{"file": t.filename, "sheet": t.sheet, "sha256": t.sha256,
                    "rows": len(t.rows)} for t in tables],
        "mappings": resolved,
        "rules_sha256": result.summary["config_sha256"],
        "csv_row_definition": "Logical record index including header; quoted multiline fields count as one record.",
    }
    return result, manifest


def csv_bytes(frame):
    """CSV opened in spreadsheet apps: prefix formula-like strings with apostrophe."""
    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return value
    return frame.map(safe).to_csv(index=False, lineterminator="\n").encode("utf-8-sig")


def report_xlsx(result, manifest):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="xlsxwriter",
                        engine_kwargs={"options": {"strings_to_formulas": False,
                                                   "strings_to_urls": False}}) as writer:
        frames = {"accepted": result.clean, "rejected": result.rejected, "changes": result.changes,
                  "summary": pd.DataFrame(result.summary.items(), columns=["metric", "value"]),
                  "inputs": pd.DataFrame(manifest["inputs"]),
                  "rules": pd.DataFrame({"json": [json.dumps(result.config, ensure_ascii=False)]}),
                  "mappings": pd.DataFrame({"table": list(manifest["mappings"]),
                      "mapping": [json.dumps(v, ensure_ascii=False) for v in manifest["mappings"].values()]})}
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes(1, 0)
            sheet.autofilter(0, 0, len(frame), len(frame.columns) - 1)
            sheet.set_column(0, len(frame.columns) - 1, 22)
    return buffer.getvalue()


def report_zip(result, manifest):
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("report.xlsx", report_xlsx(result, manifest))
        for name, frame in (("accepted", result.clean), ("rejected", result.rejected), ("changes", result.changes)):
            archive.writestr(name + ".csv", csv_bytes(frame))
        for name, value in (("summary", result.summary), ("rules", result.config), ("manifest", manifest)):
            archive.writestr(name + ".json", json.dumps(value, ensure_ascii=False, indent=2))
    return buffer.getvalue()
