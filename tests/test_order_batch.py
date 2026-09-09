from io import BytesIO
import json
from pathlib import Path
import zipfile
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from order_batch import read_tables, run_batch, suggested_mapping, report_zip, report_xlsx, csv_bytes


@pytest.fixture
def config():
    return json.loads(Path("config/orders.json").read_text(encoding="utf-8"))


def csv_table(text, name="a.csv"):
    return read_tables(name, text.encode("utf-8"))


def excel(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


def test_csv_xlsx_batch_and_export(config):
    tables = csv_table("订单号,门店,日期,金额\n001,A,2026/09/08,10\n002,A,2026-02-30,20\n")
    tables += read_tables("b.xlsx", excel([["Total","Store","Order ID","Order Date"],
                                         ["10.00","A","001","2026-09-08"],
                                         ["30","B","003","2026-09-08"]]))
    result, manifest = run_batch(tables, config)
    assert result.summary["accepted_rows"] == 2
    assert result.summary["rejected_rows"] == 2
    assert result.summary["accepted_amount_total"] == "40.00"
    assert result.clean.source_file.tolist() == ["a.csv", "b.xlsx"]
    assert result.clean.source_row.tolist() == ["2", "3"]
    report = load_workbook(BytesIO(report_xlsx(result, manifest)))
    assert report["accepted"]["A2"].value == "001"
    assert set(report.sheetnames) == {"accepted","rejected","changes","summary","inputs","rules","mappings"}
    archive = zipfile.ZipFile(BytesIO(report_zip(result, manifest)))
    assert len(archive.namelist()) == 7
    assert json.loads(archive.read("summary.json"))["accepted_rows"] == 2
    assert run_batch(tables, config)[0].summary == result.summary


@pytest.mark.parametrize("text", ["", "a,b\n", "a,a\n1,2\n", "a,b\n1,2,3\n", "a,b\n1\n"])
def test_invalid_csv(text):
    with pytest.raises(ValueError):
        csv_table(text)


def test_missing_column_stops_batch(config):
    tables = csv_table("订单号,门店,日期\n001,A,2026-09-08\n")
    with pytest.raises(ValueError, match="amount"):
        run_batch(tables, config)


def test_manual_mapping_and_extra_raw(config):
    tables = csv_table("id,shop,day,money,note\n001,A,2026-09-08,2,keep me\n")
    overrides = {tables[0].key: dict(zip(["order_id","store","date","amount"],["id","shop","day","money"]))}
    result, manifest = run_batch(tables, config, overrides)
    assert "keep me" in result.clean.raw_json.iloc[0]
    assert manifest["mappings"] == overrides


def test_formula_input_rejected():
    with pytest.raises(ValueError, match="公式"):
        read_tables("a.xlsx", excel([["a","b"], ["id", "=1+2"]]))


def test_export_does_not_execute_formula(config):
    tables = csv_table("order_id,store,date,amount\n=1+2,A,2026-09-08,1\n")
    result, manifest = run_batch(tables, config)
    sheet = load_workbook(BytesIO(report_xlsx(result, manifest)))["accepted"]
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value == "=1+2"
    assert "'=1+2" in csv_bytes(result.clean).decode("utf-8-sig")


def test_semicolon_bom(config):
    tables = read_tables("a.csv", "\ufeff订单号;门店;日期;金额\n001;A;2026-09-08;1\n".encode())
    assert run_batch(tables, config)[0].clean.order_id.tolist() == ["001"]


def test_duplicate_filename(config):
    table = csv_table("order_id,store,date,amount\n001,A,2026-09-08,1\n")
    with pytest.raises(ValueError, match="重复"):
        run_batch(table + table, config)


def test_ambiguous_alias_needs_selection(config):
    tables = csv_table("order_id,store,date,amount,金额\n001,A,2026-09-08,1,2\n")
    assert suggested_mapping(tables[0], config)["amount"] == ""
    with pytest.raises(ValueError, match="amount"):
        run_batch(tables, config)


def test_bad_excel():
    with pytest.raises(ValueError):
        read_tables("a.xlsx", b"not excel")
